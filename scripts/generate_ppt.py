from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

def create_presentation(output_path):
    prs = Presentation()
    
    # Slide 1: Title
    slide_layout = prs.slide_layouts[0] # Title Slide
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "文本引导的行人重识别\n(Text-Guided Person Re-ID)"
    subtitle.text = "利用语言监督提升特征表示学习\n\n汇报人：研究员\n日期：2023年10月"
    
    # Slide 2: Background & Motivation
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "背景与动机"
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "纯视觉 ReID 面临的挑战："
    p = tf.add_paragraph()
    p.text = "• 难分负样本 (Hard Negatives)：外观极度相似但身份不同。"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• 缺乏语义理解能力 (例如无法区分“背包”和“衣服图案”)。"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "\n解决方案：文本引导 (Text-Guided)"
    p = tf.add_paragraph()
    p.text = "• 利用自然语言描述作为辅助监督信号，引导视觉特征学习。"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• 通过 CLIP 对齐视觉与文本特征空间，增强语义判别力。"
    p.level = 1

    # Slide 3: Method Architecture
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "方法：文本引导训练流程"
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "第一阶段：多模态预训练 (对比学习)"
    p = tf.add_paragraph()
    p.text = "• 数据构建：根据属性 (颜色、类型) 生成对应的文本描述。"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• 优化目标：优化图文对比损失 (Contrastive Loss)，拉近同ID图文距离。"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "\n第二阶段：纯视觉微调 (Fine-tuning)"
    p = tf.add_paragraph()
    p.text = "• 冻结文本编码器，仅微调图像编码器 (ViT)。"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• 优化目标：ID 损失 (分类) + 三元组损失 (Triplet Loss)。"
    p.level = 1

    # Slide 4: Market1501 Results
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "实验结果：Market1501"
    
    shapes = slide.shapes
    rows = 4
    cols = 3
    left = Inches(1.0)
    top = Inches(2.0)
    width = Inches(8.0)
    height = Inches(3.0)
    
    table = shapes.add_table(rows, cols, left, top, width, height).table
    
    # Column widths
    table.columns[0].width = Inches(4.0)
    table.columns[1].width = Inches(2.0)
    table.columns[2].width = Inches(2.0)
    
    # Headers
    table.cell(0, 0).text = "模型 (Model)"
    table.cell(0, 1).text = "mAP"
    table.cell(0, 2).text = "Rank-1"
    
    # Data
    data = [
        ("Baseline (ViT-Base)", "82.4%", "90.2%"),
        ("Prior MG-ReID", "86.9%", "94.1%"),
        ("Text-Guided ReID (Ours)", "88.0%", "94.3%")
    ]
    
    for i, row_data in enumerate(data):
        table.cell(i+1, 0).text = row_data[0]
        table.cell(i+1, 1).text = row_data[1]
        table.cell(i+1, 2).text = row_data[2]
        
    slide.shapes.add_textbox(Inches(1), Inches(5.5), Inches(8), Inches(1)).text_frame.text = \
        "观察：引入文本引导后，特征的判别力得到显著提升，mAP 相比基线提高了 5.6%。"

    # Slide 5: MSMT17 Results
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "实验结果：MSMT17"
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "当前状态：训练进行中 (Epoch 4/120)"
    p = tf.add_paragraph()
    p.text = "• MSMT17 是目前规模最大、最具挑战性的 ReID 数据集。"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• 预期目标：验证模型在复杂场景下的泛化能力。"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• 待更新：训练完成后将在此处填入最终性能指标。"
    p.level = 1

    # Slide 6: Conclusion
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "总结与展望"
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "总结："
    p = tf.add_paragraph()
    p.text = "• 成功复现并改进了 Text-Guided ReID 训练流程。"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• 在 Market1501 上取得了 88.0% mAP 的优异成绩，验证了文本监督的有效性。"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "\n未来工作 (Future Work)："
    p = tf.add_paragraph()
    p.text = "• 引入更丰富的详细文本描述 (例如使用 GPT-4 或 LLaMA 生成 Caption)。"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• 在更大规模的数据集 (如 LUPerson) 上进行预训练，进一步提升泛化性。"
    p.level = 1
    
    prs.save(output_path)
    print(f"Presentation saved to {output_path}")

if __name__ == '__main__':
    output_dir = "/root/ppt_workspace"
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
    create_presentation(os.path.join(output_dir, "Text_Guided_ReID_Summary_CN.pptx"))
